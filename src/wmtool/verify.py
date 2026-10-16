# -*- coding: utf-8 -*-
"""产物校验：证明"没出页面"和"没丢内容"。

这是把需求里那两条硬约束——**不出页面**、**不丢内容**——变成可执行断言的地方，
自检和回归测试都用它。

校验项：

1. **横向不出框**：每个表格的实际总宽、每张行内图片的宽度，都不得超过其所在节的可用宽度；
2. **纵向不失真**：每张图片的高度不得超过所在节的可用高度（一图一页的硬要求）；
3. **节方向有效**：横向节确实存在且尺寸合理；
4. **内容存在性**：指定的关键词必须能在正文/表格里找到（内容没丢）。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxOpen
from docx.shared import Emu

from .geometry import A4_HEIGHT, A4_WIDTH
from .logging_setup import get_logger

log = get_logger("verify")

#: 允许的溢出余量（EMU）：排版引擎的舍入误差，2pt 以内不算越界
TOLERANCE = int(2 * 12700)


@dataclass
class SectionInfo:
    """一个节的页面信息（EMU）。"""

    index: int
    width: int
    height: int
    usable_width: int
    usable_height: int

    @property
    def landscape(self) -> bool:
        return self.width > self.height


@dataclass
class VerifyIssue:
    """一条校验问题。"""

    kind: str
    detail: str
    severity: str = "error"  # error / warning


@dataclass
class VerifyResult:
    """校验结果。"""

    ok: bool = True
    sections: list[SectionInfo] = field(default_factory=list)
    issues: list[VerifyIssue] = field(default_factory=list)
    stats: dict[str, int] = field(default_factory=dict)

    def add(self, kind: str, detail: str, severity: str = "error") -> None:
        self.issues.append(VerifyIssue(kind, detail, severity))
        if severity == "error":
            self.ok = False

    def report(self) -> str:
        lines = [f"校验结果: {'通过' if self.ok else '未通过'}"]
        lines.append(f"  统计: {self.stats}")
        for s in self.sections:
            lines.append(
                f"  节 {s.index}: {s.width / 914400:.2f}×{s.height / 914400:.2f} in"
                f" ({'横向' if s.landscape else '纵向'}) 可用宽 {s.usable_width / 914400:.2f} in"
            )
        for i in self.issues:
            lines.append(f"  [{i.severity}] {i.kind}: {i.detail}")
        return "\n".join(lines)


def _section_infos(doc) -> list[SectionInfo]:  # type: ignore[no-untyped-def]
    """把文档的每个节换算成 EMU 信息。"""
    out = []
    for i, sec in enumerate(doc.sections):
        w = int(sec.page_width or A4_WIDTH)
        h = int(sec.page_height or A4_HEIGHT)
        ml = int(sec.left_margin or 0)
        mr = int(sec.right_margin or 0)
        mt = int(sec.top_margin or 0)
        mb = int(sec.bottom_margin or 0)
        out.append(
            SectionInfo(
                index=i,
                width=w,
                height=h,
                usable_width=max(1, w - ml - mr),
                usable_height=max(1, h - mt - mb),
            )
        )
    return out


def _table_width(table) -> int:  # type: ignore[no-untyped-def]
    """表格实际总宽：优先按单元格宽度求和（Word 渲染时以此为准）。"""
    try:
        widths = []
        for row in table.rows:
            row_w = sum(int(c.width or 0) for c in row.cells)
            widths.append(row_w)
        if widths:
            return max(widths)
    except Exception:
        pass
    try:
        return sum(int(c.width or 0) for c in table.columns)
    except Exception:
        return 0


def verify_document(
    path: Path,
    *,
    must_contain: list[str] | None = None,
    min_images: int = 0,
    min_tables: int = 0,
    check_vertical: bool = True,
) -> VerifyResult:
    """对合并产物做机械校验。"""
    result = VerifyResult()
    doc = DocxOpen(str(path))
    sections = _section_infos(doc)
    result.sections = sections
    if not sections:
        result.add("结构", "文档没有任何节，文件可能损坏")
        return result

    # 每个"段落/图片/表格"归它所在的节：节是按 body 里的 sectPr 切分的。
    # 逐块判定才是有意义的——拿全文档最小的可用高度去卡所有图片，
    # 会把"规规矩矩放在纵向节里的竖图"误判成超框。
    specs = _section_resolution(doc, sections)

    # ---- 表格宽度：按所在节的可用宽度判定 ----
    table_overflow = 0
    for i, table in enumerate(doc.tables):
        w = _table_width(table)
        if w <= 0:
            continue
        limit = int(specs["tables"].get(i, _widest(sections)))
        if w > limit + TOLERANCE:
            table_overflow += 1
            result.add(
                "表格超宽",
                f"第 {i + 1} 个表格宽 {w / 914400:.2f} in，超出所在节可用宽度 "
                f"{limit / 914400:.2f} in",
            )
    result.stats["tables"] = len(doc.tables)
    result.stats["tables_overflow"] = table_overflow

    # ---- 图片尺寸：按所在节的可用宽/高判定 ----
    img_overflow_w = 0
    img_overflow_h = 0
    for i, shape in enumerate(doc.inline_shapes):
        try:
            w = int(shape.width)
            h = int(shape.height)
        except Exception:
            continue
        fallback = _tightest(sections)
        width_limit, height_limit = specs["images"].get(i, fallback)  # type: ignore[misc]
        width_limit = int(width_limit)
        height_limit = int(height_limit)
        if w > width_limit + TOLERANCE:
            img_overflow_w += 1
            result.add(
                "图片超宽",
                f"第 {i + 1} 张图宽 {w / 914400:.2f} in，超出所在节可用宽度 "
                f"{width_limit / 914400:.2f} in",
            )
        if check_vertical and h > height_limit + TOLERANCE:
            img_overflow_h += 1
            result.add(
                "图片超高",
                f"第 {i + 1} 张图高 {h / 914400:.2f} in，超出所在节可用高度 "
                f"{height_limit / 914400:.2f} in（一图一页要求图片完整落在一页内）",
            )
    result.stats["images"] = len(doc.inline_shapes)
    result.stats["images_overflow_w"] = img_overflow_w
    result.stats["images_overflow_h"] = img_overflow_h

    # ---- 节方向 ----
    landscape = [s for s in sections if s.landscape]
    result.stats["sections"] = len(sections)
    result.stats["landscape_sections"] = len(landscape)
    for s in sections:
        if s.width < int(4 * 914400) or s.height < int(4 * 914400):
            result.add("节尺寸异常", f"节 {s.index} 尺寸过小: {s.width}×{s.height} EMU")

    # ---- 内容存在性 ----
    if must_contain:
        text = _full_text(doc)
        missing = [k for k in must_contain if k not in text]
        if missing:
            result.add("内容缺失", "以下关键词未出现在产物里: " + "、".join(missing))
        result.stats["keywords_checked"] = len(must_contain)

    if min_images and len(doc.inline_shapes) < min_images:
        result.add("图片数不足", f"期望至少 {min_images} 张图，实际 {len(doc.inline_shapes)} 张")
    if min_tables and len(doc.tables) < min_tables:
        result.add("表格数不足", f"期望至少 {min_tables} 个表，实际 {len(doc.tables)} 个")

    result.stats["paragraphs"] = len(doc.paragraphs)
    return result


def _full_text(doc) -> str:  # type: ignore[no-untyped-def]
    """把正文、表格、页眉页脚的文字拼起来（用于关键词检查）。"""
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for sec in doc.sections:
        for hf in (sec.header, sec.footer):
            try:
                parts.extend(p.text for p in hf.paragraphs)
            except Exception:
                pass
    return "\n".join(parts)


def _widest(sections: list[SectionInfo]) -> int:
    """最宽的可用宽度（兜底用）。"""
    return max(s.usable_width for s in sections)


def _tightest(sections: list[SectionInfo]) -> tuple[int, int]:
    """最紧的可用宽/高（兜底用）。"""
    return min(s.usable_width for s in sections), min(s.usable_height for s in sections)


def _section_resolution(doc, sections: list[SectionInfo]) -> dict[str, dict[int, object]]:  # type: ignore[no-untyped-def]
    """算出每个表格/图片落在哪个节。

    OOXML 的节划分有个坑：``w:sectPr`` **通常不是 body 的直接子元素**，
    而是挂在某个段落 ``w:pPr`` 里的（python-docx 的 ``add_section`` 就是这么写的），
    最后一个节才以 body 直属 ``w:sectPr`` 收尾。

    所以正确算法是：按顺序遍历 body 的段落/表格，
    **遇到"带 sectPr 的段落"才算跨进下一个节**（该段落本身仍属旧节），
    元素归它当时所在的节。
    """
    from docx.oxml.ns import qn

    body = doc.element.body
    table_sec: dict[int, int] = {}
    image_sec: dict[int, int] = {}

    sec_index = 0
    t_idx = 0
    img_idx = 0
    last_sec = max(0, len(sections) - 1)

    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            ppr = child.find(qn("w:pPr"))
            has_sect = ppr is not None and ppr.find(qn("w:sectPr")) is not None
            # 先登记本段里的图片（属于当前节），再处理节切换
            for _ in range(len(child.findall(".//" + qn("w:drawing")))):
                image_sec[img_idx] = min(sec_index, last_sec)
                img_idx += 1
            if has_sect:
                sec_index += 1
        elif child.tag == qn("w:tbl"):
            table_sec[t_idx] = min(sec_index, last_sec)
            t_idx += 1

    def _limits(sec_i: int) -> tuple[int, int]:
        s = sections[min(sec_i, last_sec)]
        return s.usable_width, s.usable_height

    return {
        "tables": {i: _limits(si)[0] for i, si in table_sec.items()},
        "images": {i: _limits(si) for i, si in image_sec.items()},
    }

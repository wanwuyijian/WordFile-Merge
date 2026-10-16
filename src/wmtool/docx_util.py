# -*- coding: utf-8 -*-
"""文档构建辅助：分节、分页、横向页、图片插入、书签、目录域。

python-docx 对这些东西的支持要么没有（书签、横向节），要么很绕（域代码），
所以统一收在这一层，转换器只管调用。
"""

from __future__ import annotations

from typing import Iterable, Sequence

from docx.document import Document as DocxDocument
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor
from docx.table import Table

from .geometry import A4_HEIGHT, A4_WIDTH, Margins, PageBox
from .logging_setup import get_logger

log = get_logger("docx_util")


# --------------------------------------------------------------------------- #
# 页面 / 分节
# --------------------------------------------------------------------------- #
def configure_section(section, box: PageBox) -> None:  # type: ignore[no-untyped-def]
    """把某个 section 设成指定的页面框（尺寸 + 方向 + 页边距）。"""
    landscape = box.width > box.height
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Emu(int(box.width))
    section.page_height = Emu(int(box.height))
    m = box.margins
    section.top_margin = Emu(m.top)
    section.bottom_margin = Emu(m.bottom)
    section.left_margin = Emu(m.left)
    section.right_margin = Emu(m.right)


def new_section(
    doc: DocxDocument,
    box: PageBox,
    *,
    new_page: bool = True,
    restart_page_number: bool = False,
):
    """新建一个 section 并返回它。

    ``restart_page_number=True`` 用于"一节一页"的横向内容——
    页码重新从 1 开始，避免 Word 里横向页把全局页码带偏。
    """
    start_type = WD_SECTION.NEW_PAGE if new_page else WD_SECTION.CONTINUOUS
    section = doc.add_section(start_type)
    configure_section(section, box)
    if restart_page_number:
        _set_page_numbering(section, start=1)
    return section


def _set_page_numbering(section, *, start: int | None = None, fmt: str | None = None) -> None:  # type: ignore[no-untyped-def]
    """设置该 section 的页码起始值 / 编号格式。"""
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    if start is not None:
        pg_num.set(qn("w:start"), str(start))
    if fmt is not None:
        pg_num.set(qn("w:fmt"), fmt)


def set_document_defaults(doc: DocxDocument, box: PageBox = PageBox(width=A4_WIDTH, height=A4_HEIGHT)) -> None:
    """设置文档首个 section 与默认样式（中文字体、字号、行距）。"""
    configure_section(doc.sections[0], box)

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(10.5)
    # 中文字体必须同时设置 eastAsia，否则 Word 里显示的是宋体
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "微软雅黑")
    rfonts.set(qn("w:hAnsi"), "微软雅黑")
    rfonts.set(qn("w:eastAsia"), "微软雅黑")

    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.15


def page_break(doc: DocxDocument) -> None:
    """插入一个硬分页符。"""
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_heading(doc: DocxDocument, text: str, *, level: int = 1, size_pt: float | None = None,
                page_break_before: bool = False) -> None:
    """加一个标题段落（不用内置 Heading 样式，避免踩到模板样式缺失）。"""
    p = doc.add_paragraph()
    if page_break_before:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(10 if level <= 1 else 6)
    p.paragraph_format.space_after = Pt(6 if level <= 1 else 4)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "微软雅黑"
    run.font.size = Pt(size_pt if size_pt is not None else (16 if level == 1 else 13))
    run.font.color.rgb = RGBColor(0x1F, 0x36, 0x5D)
    _set_run_east_asian(run, "微软雅黑")


def _set_run_east_asian(run, font_name: str) -> None:  # type: ignore[no-untyped-def]
    """给 run 设置中文字体（python-docx 的 run.font.name 不管 eastAsia）。"""
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)


def add_caption(doc: DocxDocument, text: str) -> None:
    """加一行灰色小字说明（源文件名、页码之类）。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    _set_run_east_asian(run, "微软雅黑")


# --------------------------------------------------------------------------- #
# 图片
# --------------------------------------------------------------------------- #
def add_picture_centered(
    doc: DocxDocument,
    stream,
    width_emu: int,
    height_emu: int,
    *,
    alt: str = "",
) -> None:
    """居中插入一张已算好尺寸的图片。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    pic = run.add_picture(stream, width=Emu(int(width_emu)), height=Emu(int(height_emu)))
    if alt:
        try:
            doc_pr = pic._inline.docPr
            doc_pr.set("descr", alt)
            doc_pr.set("title", alt)
        except Exception:
            pass


# --------------------------------------------------------------------------- #
# 表格
# --------------------------------------------------------------------------- #
def style_data_table(table: Table, *, font_pt: float = 9.0) -> None:
    """给数据表套上边框 + 表头底色，并把全文统一到指定字号。"""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(font_pt)
                    _set_run_east_asian(run, "微软雅黑")
        if r_idx == 0:
            for cell in row.cells:
                _shade_cell(cell, "DCE6F1")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True


def _set_table_borders(table: Table, *, sz: int = 6, color: str = "808080") -> None:
    """给表格加全边框（python-docx 没有现成 API）。"""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def _shade_cell(cell, fill: str) -> None:  # type: ignore[no-untyped-def]
    """给单元格加底色。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_table_layout_fixed(table: Table) -> None:
    """固定表格布局，让列宽完全按我们算的值走，Word 不再自作主张重排。"""
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def repeat_header_row(table: Table) -> None:
    """把首行设为跨页重复的表头。"""
    if not table.rows:
        return
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_text(cell, text: str, *, bold: bool = False, size_pt: float = 9.0,
                  align: WD_ALIGN_PARAGRAPH | None = None) -> None:  # type: ignore[no-untyped-def]
    """写单元格文本，保留换行，统一字体。"""
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    first = True
    for line in str(text).split("\n"):
        if not first:
            p.add_run().add_break()
        run = p.add_run(line)
        run.bold = bold
        run.font.size = Pt(size_pt)
        _set_run_east_asian(run, "微软雅黑")
        first = False


def set_row_height_pt(row, pt: float) -> None:  # type: ignore[no-untyped-def]
    """设行高（至少 pt 磅）。"""
    tr_pr = row._tr.get_or_add_trPr()
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), str(int(pt * 20)))  # 磅 -> 缇
    height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(height)


# --------------------------------------------------------------------------- #
# 书签 / 目录域
# --------------------------------------------------------------------------- #
_BOOKMARK_ID = 0


def _next_bookmark_id() -> int:
    global _BOOKMARK_ID
    _BOOKMARK_ID += 1
    return _BOOKMARK_ID


def add_bookmark(doc: DocxDocument, name: str) -> None:
    """在当前位置插一个书签（Word 导航窗格 / 超链接可跳转）。

    书签名必须符合 Word 规则：以字母开头、只含字母数字下划线、不超过 40 字符。
    """
    safe = "".join(ch if ch.isalnum() or ch == "_" else "_" for ch in name)
    if not safe or not safe[0].isalpha():
        safe = "bm_" + safe
    safe = safe[:40]
    bid = _next_bookmark_id()

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), safe)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    p._p.append(start)
    p._p.append(end)


def add_toc_field(doc: DocxDocument, levels: str = "1-3") -> None:
    """插入 Word 目录域。

    域的结果本身是缓存值：Word 打开后按 F9（或"更新目录"）即可展开真实页码。
    这里额外塞一段提示文字，避免用户看到空白以为坏了。
    """
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f' TOC \\o "{levels}" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "（此处为自动目录：在 Word 中按 Ctrl+A 再按 F9，或右键选择“更新域”即可生成页码）"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, fld_text, fld_end):
        run._r.append(el)

    # 打开文档时自动提示更新域
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_hyperlink(paragraph, url: str, text: str) -> None:  # type: ignore[no-untyped-def]
    """在段落里插一个外部超链接。"""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def join_texts(parts: Iterable[str], sep: str = "；") -> str:
    """拼接非空文本片段。"""
    return sep.join(p for p in (str(x).strip() for x in parts) if p)


def merge_cells_safe(row, start: int, end: int) -> None:  # type: ignore[no-untyped-def]
    """安全合并单元格（越界/单格时静默跳过）。"""
    try:
        if end > start and len(row.cells) > end:
            row.cells[start].merge(row.cells[end])
    except Exception as exc:
        log.debug("合并单元格失败 (%d-%d): %s", start, end, exc)


def style_names(doc: DocxDocument) -> Sequence[str]:
    """调试用：列出文档可用样式名。"""
    return [s.name for s in doc.styles]
